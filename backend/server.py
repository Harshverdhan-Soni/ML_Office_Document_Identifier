from fastapi import FastAPI, File, UploadFile
from fastapi.middleware.cors import CORSMiddleware
import shutil, os
from PyPDF2 import PdfReader
import docx
import pytesseract
from PIL import Image
import openpyxl
import joblib  # load your trained pipeline

# === Load trained pipeline (vectorizer + classifier) ===
MODEL_PATH = "model.pkl"  # your saved pipeline
ml_pipeline = joblib.load(MODEL_PATH)

# === FastAPI setup ===
app = FastAPI()
UPLOAD_FOLDER = "uploads"
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

# Enable CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# === Text extraction function ===
def extract_text(file_path: str) -> str:
    ext = file_path.split('.')[-1].lower()
    text = ""
    if ext == "pdf":
        reader = PdfReader(file_path)
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext in ["doc", "docx"]:
        doc = docx.Document(file_path)
        text = "\n".join([p.text for p in doc.paragraphs])
    elif ext in ["xls", "xlsx"]:
        wb = openpyxl.load_workbook(file_path)
        sheet = wb.active
        text = "\n".join(
            [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value]
        )
    elif ext in ["jpg", "jpeg", "png"]:
        text = pytesseract.image_to_string(Image.open(file_path))
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return text

# === ML classification function ===
def classify_document_from_text(text: str) -> str:
    # ml_pipeline already includes vectorizer + classifier
    predicted_label = ml_pipeline.predict([text])[0]  # wrap text in list
    return predicted_label

# === FastAPI endpoint ===
@app.post("/classify")
async def classify(file: UploadFile = File(...)):
    file_path = os.path.join(UPLOAD_FOLDER, file.filename)
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)

    try:
        text = extract_text(file_path)
        doc_type = classify_document_from_text(text)
    except Exception as e:
        doc_type = f"Error: {str(e)}"

    # Remove temp file
    os.remove(file_path)
    return {"filename": file.filename, "extracted": text, "type": doc_type}
